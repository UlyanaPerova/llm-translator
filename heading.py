#!/usr/bin/env python3
"""
heading.py — Форматирование заголовков глав в .docx

Функции:
  1. Дедупликация повторяющихся заголовков глав
     (Глава N / Глава N / Глава N. Название → оставляет самый полный)
  2. Назначение стиля Heading 1 + разрыв страницы перед каждой главой
  3. Сохранение бэкапа оригинала перед изменениями

Использование:
  python3 heading.py input.docx [-o output.docx]
"""

import argparse
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

from logger import setup_logger
import logging

setup_logger(prefix="heading")
log = logging.getLogger("heading")

CHAPTER_RE = re.compile(r"^(?:Глава|Chapter)\s+(\d+)", re.IGNORECASE)


def dedup_chapter_headings(doc: Document) -> int:
    """Убирает дубли заголовков глав, оставляя самый длинный (с названием)."""
    paragraphs = list(doc.paragraphs)
    to_remove = set()
    removed = 0
    i = 0

    while i < len(paragraphs):
        m = CHAPTER_RE.match(paragraphs[i].text.strip())
        if m:
            num = m.group(1)
            group = [i]
            j = i + 1
            while j < len(paragraphs):
                mj = CHAPTER_RE.match(paragraphs[j].text.strip())
                if mj and mj.group(1) == num:
                    group.append(j)
                    j += 1
                else:
                    break

            if len(group) > 1:
                best = max(group, key=lambda idx: len(paragraphs[idx].text.strip()))
                for idx in group:
                    if idx != best:
                        to_remove.add(idx)
                        removed += 1
                log.debug(
                    "Глава %s: %d дублей → оставлено «%s»",
                    num, len(group) - 1, paragraphs[best].text.strip()[:60],
                )
            i = j
        else:
            i += 1

    for idx in sorted(to_remove, reverse=True):
        el = paragraphs[idx]._element
        el.getparent().remove(el)

    return removed


def apply_heading_style(doc: Document) -> int:
    """Назначает Heading 1 + разрыв страницы для заголовков глав."""
    count = 0
    for para in doc.paragraphs:
        if CHAPTER_RE.match(para.text.strip()):
            para.style = doc.styles["Heading 1"]
            pPr = para._p.get_or_add_pPr()
            page_break = OxmlElement("w:pageBreakBefore")
            pPr.append(page_break)
            count += 1
    return count


def main():
    parser = argparse.ArgumentParser(
        description="Форматирование заголовков глав в .docx (дедупликация + Heading 1)"
    )
    parser.add_argument("input", help="Входной .docx файл")
    parser.add_argument("-o", "--output", help="Выходной файл (по умолчанию: input_heading.docx)")
    parser.add_argument("--no-dedup", action="store_true", help="Не убирать дубли заголовков")
    parser.add_argument("--no-backup", action="store_true", help="Не создавать бэкап")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        log.error("Файл не найден: %s", input_path)
        sys.exit(1)

    if args.output:
        output_path = Path(args.output)
    else:
        output_path = input_path.with_name(f"{input_path.stem}_heading{input_path.suffix}")

    # Бэкап
    if not args.no_backup:
        backup_dir = Path("backups")
        backup_dir.mkdir(exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = backup_dir / f"{input_path.stem}_{timestamp}{input_path.suffix}"
        shutil.copy2(input_path, backup_path)
        log.info("Бэкап: %s", backup_path)

    log.info("Открываю: %s", input_path)
    doc = Document(str(input_path))

    # Дедупликация
    if not args.no_dedup:
        removed = dedup_chapter_headings(doc)
        log.info("Убрано дублей заголовков: %d", removed)

    # Форматирование
    styled = apply_heading_style(doc)
    log.info("Отформатировано заголовков: %d", styled)

    doc.save(str(output_path))
    log.info("Сохранено: %s", output_path)


if __name__ == "__main__":
    main()
