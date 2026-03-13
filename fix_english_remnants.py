"""
Быстрый фикс: находит оставшиеся английские фразы в переведённом .docx
и переводит их через Gemini 2.5 Flash (дёшево).

Использование:
    python3 fix_english_remnants.py nepravilno_ponyaty_okhotnik_iz_drugogo_mira_p_1_380_415.docx
    python3 fix_english_remnants.py file.docx -o file_fixed.docx
    python3 fix_english_remnants.py file.docx --dry-run   # только показать что найдёт
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path
from copy import deepcopy

from dotenv import load_dotenv
from docx import Document
from logger import setup_logger
import logging

load_dotenv()

setup_logger(prefix="fix_english")
log = logging.getLogger("fix_english")

# Слова, которые нормально оставлять на английском
WHITELIST = {
    # Бренды и общеупотребительные
    "ok", "vip", "tv", "youtube", "google", "iphone", "android", "wifi",
    "sms", "gps", "id", "pc", "npc", "hp", "mp", "xp", "pvp", "pve",
    "rpg", "mmorpg", "atk", "def", "dps", "aoe", "buff", "debuff",
    "login", "online", "offline", "live", "vs", "ui", "ai",
    # Ранги / классы
    "ss", "sss", "ex",
    # Единицы / аббревиатуры
    "kg", "km", "cm", "mm", "pm", "am",
    # Файлы / техническое
    "pdf", "jpg", "png", "gif", "url", "http", "https",
    # HTML-артефакты / фрагменты
    "gt", "lt", "amp", "fi", "wi",
    # Общеупотребительные в русском
    "sos", "ceo", "iq", "no", "yes", "the", "pp", "oo", "ww",
    # Страны
    "uk", "usa",
    # Междометия
    "ha", "haha", "hmm", "oh", "ah", "ugh", "tsk", "pfft", "heh",
    "hoo", "whoa", "wow", "ooh", "aah", "eh", "uh", "um",
    # Форматирование / служебные
    "chapter", "vol",
}

# Паттерн: 2+ подряд латинских буквы (слово), исключая HTML-теги
ENGLISH_WORD_RE = re.compile(r"\b[A-Za-z]{2,}\b")
HTML_TAG_RE = re.compile(r"</?[bi]>")


def has_english(text: str) -> list[str]:
    """Находит английские слова в тексте (игнорирует теги и whitelist)."""
    clean = HTML_TAG_RE.sub("", text)
    words = ENGLISH_WORD_RE.findall(clean)
    return [w for w in words if w.lower() not in WHITELIST]


def fix_paragraph_gemini(client, text: str, eng_words: list[str]) -> str:
    """Отправляет абзац в Gemini для замены английских слов русскими."""
    from google.genai import types

    prompt = (
        "В этом русском тексте остались непереведённые английские слова/фразы. "
        "Переведи или транслитерируй их на русский, сохраняя контекст и падеж. "
        "Верни ТОЛЬКО исправленный текст, без комментариев. "
        "Сохрани все HTML-теги (<b>, </b>, <i>, </i>) на местах.\n\n"
        f"Непереведённые слова: {', '.join(set(eng_words))}\n\n"
        f"Текст:\n{text}"
    )

    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=[prompt],
        config=types.GenerateContentConfig(
            temperature=0,
            max_output_tokens=4000,
        ),
    )

    if response.text is None:
        log.warning("Gemini вернул None — оставляю как есть")
        return text

    return response.text.strip()


def main():
    parser = argparse.ArgumentParser(
        description="Фикс английских слов в переведённом .docx через Gemini"
    )
    parser.add_argument("input", help="Путь к переведённому .docx файлу")
    parser.add_argument(
        "-o", "--output", default=None,
        help="Выходной файл (по умолчанию: {имя}_fixed.docx)",
    )
    parser.add_argument(
        "--dry-run", action="store_true",
        help="Только показать найденные английские слова, без исправления",
    )

    args = parser.parse_args()

    if not os.path.isfile(args.input):
        sys.exit(f"Файл не найден: {args.input}")

    # Выходной путь
    if args.output:
        output_path = args.output
    else:
        stem = Path(args.input).stem
        output_path = f"{stem}_fixed.docx"

    # Читаем документ
    doc = Document(args.input)
    paragraphs_to_fix = []

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        eng_words = has_english(text)
        if eng_words:
            paragraphs_to_fix.append((i, para, eng_words))

    print(f"Найдено {len(paragraphs_to_fix)} абзацев с английскими словами")

    if not paragraphs_to_fix:
        print("Нечего исправлять!")
        return

    # Показать что нашли
    all_eng = set()
    for i, para, eng_words in paragraphs_to_fix:
        all_eng.update(w.lower() for w in eng_words)

    print(f"\nУникальные английские слова ({len(all_eng)}):")
    for w in sorted(all_eng):
        print(f"  - {w}")

    if args.dry_run:
        print(f"\n--dry-run: показано {len(paragraphs_to_fix)} абзацев, без исправлений")
        return

    # Инициализируем Gemini
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        sys.exit("GEMINI_API_KEY не найден в .env")

    from google import genai
    client = genai.Client(api_key=api_key)

    # Исправляем
    fixed_count = 0
    for idx, (i, para, eng_words) in enumerate(paragraphs_to_fix):
        # Собираем текст абзаца с форматированием (через runs)
        original_text = ""
        for run in para.runs:
            t = run.text
            if run.bold and run.italic:
                original_text += f"<b><i>{t}</i></b>"
            elif run.bold:
                original_text += f"<b>{t}</b>"
            elif run.italic:
                original_text += f"<i>{t}</i>"
            else:
                original_text += t

        print(f"\n[{idx+1}/{len(paragraphs_to_fix)}] Английские: {', '.join(eng_words)}")
        log.info("[%d/%d] Абзац %d: %s", idx+1, len(paragraphs_to_fix), i, eng_words)

        fixed_text = fix_paragraph_gemini(client, original_text, eng_words)

        # Проверяем, что что-то изменилось
        if fixed_text == original_text:
            print("  → без изменений")
            continue

        remaining = has_english(fixed_text)
        if remaining:
            log.info("  Осталось: %s", remaining)

        # Заменяем текст в абзаце — очищаем все runs и пишем заново
        # Сохраняем форматирование через HTML-теги
        from ocr_vision import _parse_formatting

        # Удаляем старые runs
        for run in para.runs:
            run.text = ""

        # Парсим новый текст и создаём runs
        # Сначала удалим все runs кроме первого
        while len(para.runs) > 1:
            p_element = para._p
            p_element.remove(para.runs[-1]._r)

        segments = _parse_formatting(fixed_text)
        first = True
        for seg_text, bold, italic in segments:
            if first and para.runs:
                run = para.runs[0]
                run.text = seg_text
                run.bold = bold if bold else None
                run.italic = italic if italic else None
                first = False
            else:
                run = para.add_run(seg_text)
                run.bold = bold if bold else None
                run.italic = italic if italic else None

        fixed_count += 1
        print(f"  → исправлено")

    # Сохраняем
    doc.save(output_path)
    print(f"\nГотово: {fixed_count} абзацев исправлено → {output_path}")


if __name__ == "__main__":
    main()
