#!/usr/bin/env python3
"""
Ремонт pro-перевода: переводит только упавшие чанки (из failed_chunks_pro.json)
и заменяет в docx маркеры «[ОШИБКА ПЕРЕВОДА ЧАНКА N: ...]» на свежий перевод.

Успешные 999 pro-чанков не переоплачиваются. Прогресс кэшируется — обрыв
(в т.ч. по дневной квоте) не теряет ни денег, ни работы: перезапуск продолжит.

Использование:
    python3 repair_pro_translation.py                # всё по умолчанию
    python3 repair_pro_translation.py --parallel 8
"""

import argparse
import json
import os
import re
import sys
import logging
from pathlib import Path
from dotenv import load_dotenv

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # корень проекта в sys.path
from common.logger import setup_logger
setup_logger(prefix="repair_pro")
log = logging.getLogger("repair_pro")

load_dotenv()

from docx import Document
from translators.eng_translator import (
    extract_text, split_into_chunks, load_glossary, _strip_html_tags, _parse_formatting,
)
from docx.shared import Pt, Cm
import translators.eng_translator_gemini as tg

SOURCE = "The_Artist_Who_Paints_Dungeon.docx"
TARGET = "The_Artist_Who_Paints_Dungeon_translated_pro.docx"
GLOSSARY = "The_Artist_Who_Paints_Dungeon_glossary.json"
FAILED_LIST = "failed_chunks_pro.json"
REPAIR_CACHE = ".repair_pro_cache.json"
ERROR_RE = re.compile(r"\[ОШИБКА ПЕРЕВОДА ЧАНКА (\d+):")


def load_repair_cache() -> dict[int, str]:
    if os.path.isfile(REPAIR_CACHE):
        with open(REPAIR_CACHE, "r", encoding="utf-8") as f:
            return {int(k): v for k, v in json.load(f).items()}
    return {}


def save_repair_cache(done: dict[int, str]) -> None:
    tmp = REPAIR_CACHE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump({str(k): v for k, v in done.items()}, f, ensure_ascii=False)
    os.replace(tmp, REPAIR_CACHE)


def insert_translation_before(doc, error_para, chunk_text: str) -> None:
    """Вставляет абзацы перевода перед error_para (стиль как у save_to_docx)."""
    for block in chunk_text.split("\n\n"):
        for line in block.split("\n"):
            line = line.strip()
            if not line:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.first_line_indent = Cm(1.25)
            for seg_text, bold, italic in _parse_formatting(line):
                run = p.add_run(seg_text)
                if bold:
                    run.bold = True
                if italic:
                    run.italic = True
            error_para._p.addprevious(p._p)


def main():
    parser = argparse.ArgumentParser(description="Ремонт pro-перевода (только упавшие чанки)")
    parser.add_argument("--parallel", type=int, default=8)
    parser.add_argument("--model", default="gemini-2.5-pro")
    args = parser.parse_args()

    failed = json.load(open(FAILED_LIST))
    print(f"Чанков к ремонту: {len(failed)}")

    text = extract_text(SOURCE)
    chunks = split_into_chunks(text, 3000)
    glossary = load_glossary(GLOSSARY)
    relations = tg.load_relations(GLOSSARY)
    litrpg = tg.load_litrpg_templates()
    pool = tg.make_client_pool()
    print(f"API-ключей в пуле: {len(pool)}")

    # Конспекты для контекста — дешёвой моделью (flash, у неё квота свободна)
    cached = tg.load_cache(".repair_summaries_cache.json", len(chunks))
    summaries = tg.build_summaries(pool, chunks, cached.get("summaries"), workers=8)
    tg.save_cache(".repair_summaries_cache.json",
                  {"total_chunks": len(chunks), "summaries": summaries, "results": {}})

    done = load_repair_cache()
    todo = [i for i in failed if i not in done]
    print(f"Уже в кэше: {len(done)}, осталось перевести: {len(todo)}")

    # Перевод упавших чанков (параллельно)
    import concurrent.futures, threading
    lock = threading.Lock()

    def _work(i: int) -> None:
        ctx = tg.build_story_context(summaries, i - 1)
        result = tg.translate_chunk(
            pool, args.model, chunks[i - 1], i, len(chunks),
            glossary=glossary, relations=relations, litrpg=litrpg,
            story_context=ctx, quiet=True,
        )
        if result.startswith("[ОШИБКА"):
            # НЕ кэшируем ошибки (кончились деньги/сбой) — чанк останется в todo,
            # а в docx сохранится маркер для следующего захода ремонта
            log.warning("Чанк %d не переведён (ошибка), останется на потом", i)
            print(f"  ❌ чанк {i} — отложен")
            return
        with lock:
            done[i] = result
            save_repair_cache(done)
            print(f"  ✅ чанк {i} ({len(done)}/{len(failed)})")

    stopped_early = None
    try:
        with concurrent.futures.ThreadPoolExecutor(max_workers=args.parallel) as ex:
            list(ex.map(_work, todo))
    except tg.DailyQuotaExceeded as e:
        save_repair_cache(done)
        stopped_early = str(e)
        print(f"\n🛑 {e}")
        print(f"   Прогресс: {len(done)}/{len(failed)}. Вживляю в docx то, что готово;")
        print(f"   остальное — следующим заходом (перезапусти скрипт позже).")

    # Вживление в docx
    print("Вживляю переводы в docx...")
    import shutil
    from datetime import datetime
    backup = f"{Path(TARGET).stem}_before_repair_{datetime.now():%Y%m%d_%H%M%S}.docx"
    shutil.copy2(TARGET, backup)
    print(f"Бэкап: {backup}")

    doc = Document(TARGET)
    replaced = 0
    for para in list(doc.paragraphs):
        m = ERROR_RE.match(para.text.strip())
        if not m:
            continue
        idx = int(m.group(1))
        if idx not in done:
            log.warning("Чанк %d не переведён, маркер оставлен", idx)
            continue
        insert_translation_before(doc, para, done[idx])
        para._p.getparent().remove(para._p)
        replaced += 1
    doc.save(TARGET)
    print(f"Заменено маркеров: {replaced}/{len(failed)}")

    # Финальная проверка
    doc = Document(TARGET)
    full = "\n".join(p.text for p in doc.paragraphs)
    leftover_errors = len(ERROR_RE.findall(full))
    print(f"\nИтог: {len(doc.paragraphs):,} абзацев, {len(full):,} символов")
    print(f"Оставшихся маркеров ошибок: {leftover_errors}")
    print(f"💰 Стоимость ремонта: ${tg.actual_cost():.2f}")
    if leftover_errors == 0:
        os.remove(REPAIR_CACHE)
        print("Кэш ремонта удалён. Готово ✅")
    elif stopped_early:
        print(f"Частичная сборка: {leftover_errors} чанков ждут следующего захода "
              f"(причина остановки: {stopped_early})")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nПрервано. Перезапуск продолжит с кэша.")
        sys.exit(1)
