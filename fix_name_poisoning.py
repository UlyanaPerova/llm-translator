#!/usr/bin/env python3
"""
Точечный фикс отравленных имён в переводе (баг кросс-алиасов глоссария).

Чинит ТОЛЬКО:
  - «Зеордж/Зеорге/Джордж» (любые падежные формы) → «Джио» / «Со Джио»
  - «Улыбающийся тиран морской деревни» (любые формы) → «Джованни»

Ничего другого не меняет: каждый ответ GPT проверяется пословным диффом;
если модель тронула что-то ещё — её версия отбрасывается и применяется
детерминированная замена. Прогресс кэшируется, перезапуск продолжает.
"""

import difflib
import json
import os
import re
import sys
import logging
import shutil
from datetime import datetime

from dotenv import load_dotenv
load_dotenv()

from logger import setup_logger
setup_logger(prefix="fix_names")
log = logging.getLogger("fix_names")

from docx import Document
from openai import OpenAI
from eng_translator import _parse_formatting

TARGET = "The_Artist_Who_Paints_Dungeon_translated_pro.docx"
CACHE = ".fix_names_cache.json"
MODEL = "gpt-5.1"

BAD_NAME_RE = re.compile(r"(?:Зеордж|Зеорге|Джордж)[а-яё]*")
TYRANT_RE = re.compile(r"[Уу]лыбающ[а-яё]+\s+тиран[а-яё]*\s+морской\s+деревни")
ANY_BAD_RE = re.compile(rf"(?:{BAD_NAME_RE.pattern})|(?:{TYRANT_RE.pattern})")

# слова, которые допустимо УДАЛЯТЬ (части испорченных имён)
REMOVABLE_RE = re.compile(
    r"^[«»\"'(\[]?([Зз]еордж|[Зз]еорге|[Дд]жордж|[Уу]лыбающ|[Тт]иран|морской|деревни)",
)
# слова, которые допустимо ВСТАВЛЯТЬ (правильные имена)
INSERTABLE_RE = re.compile(r"^[«»\"'(\[]?(Джио|Со|Джованни)[.,!?;:…»\"')\]]*$")

SYSTEM_PROMPT = """You are a precise text-repair tool for a Russian novel translation.
A glossary bug corrupted one character's name. In the paragraph you receive:

1. «Зеордж», «Зеорге», «Джордж» (any grammatical form: «Зеорджа», «мистером Зеорджем»...)
   are corrupted renderings of the SAME character. His correct Russian names:
   «Джио» (short form) or «Со Джио» (full form: «мистер Со Джио», «охотник Со Джио»).
   Replace each corrupted occurrence with whichever correct form reads naturally
   in context. Both forms are indeclinable — use them unchanged in every case.
2. The phrase «Улыбающийся тиран морской деревни» (any grammatical form) must be
   replaced by the name «Джованни» (indeclinable).

Change NOTHING else: not a single other word, punctuation mark, or <b>/<i> tag.
Return ONLY the corrected paragraph text."""


def deterministic_fix(text: str) -> str:
    text = TYRANT_RE.sub("Джованни", text)
    text = BAD_NAME_RE.sub("Джио", text)
    return text


def diff_is_safe(original: str, fixed: str) -> bool:
    """True, если изменения затронули только испорченные имена."""
    a, b = original.split(), fixed.split()
    sm = difflib.SequenceMatcher(a=a, b=b, autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        for w in a[i1:i2]:
            if not REMOVABLE_RE.match(w):
                return False
        for w in b[j1:j2]:
            if not INSERTABLE_RE.match(w):
                return False
    return True


def para_tagged_text(para) -> str:
    """Текст абзаца с <b>/<i>-тегами из runs (как в fix_english_remnants)."""
    out = ""
    for run in para.runs:
        t = run.text
        if not t:
            continue
        if run.bold and run.italic:
            out += f"<b><i>{t}</i></b>"
        elif run.bold:
            out += f"<b>{t}</b>"
        elif run.italic:
            out += f"<i>{t}</i>"
        else:
            out += t
    return out or para.text


def rewrite_para(para, new_text: str) -> None:
    """Переписывает runs абзаца новым текстом, сохраняя <b>/<i>."""
    while len(para.runs) > 1:
        para._p.remove(para.runs[-1]._r)
    segments = _parse_formatting(new_text)
    first = True
    for seg_text, bold, italic in segments:
        if first and para.runs:
            run = para.runs[0]
            run.text = seg_text
        else:
            run = para.add_run(seg_text)
        run.bold = bold if bold else None
        run.italic = italic if italic else None
        first = False


def main():
    api_key = os.getenv("OPENAI_API_KEY") or sys.exit("OPENAI_API_KEY не найден")
    client = OpenAI(api_key=api_key)

    doc = Document(TARGET)
    paras = doc.paragraphs

    affected = [i for i, p in enumerate(paras) if ANY_BAD_RE.search(p.text)]
    print(f"затронутых абзацев: {len(affected)}")
    if not affected:
        print("нечего чинить")
        return

    cache: dict[str, str] = {}
    if os.path.isfile(CACHE):
        cache = json.load(open(CACHE, encoding="utf-8"))
        print(f"из кэша: {len(cache)}")

    # Фаза 1: детерминированно чиним абзацы, где ТОЛЬКО фраза-эпитет (без Зеорджей)
    det_only = [i for i in affected
                if TYRANT_RE.search(paras[i].text) and not BAD_NAME_RE.search(paras[i].text)]
    for i in det_only:
        if str(i) not in cache:
            cache[str(i)] = deterministic_fix(para_tagged_text(paras[i]))
    print(f"детерминированно (только эпитет): {len(det_only)}")

    # Фаза 2: GPT для абзацев с испорченными именами
    gpt_todo = [i for i in affected if str(i) not in cache]
    print(f"через GPT: {len(gpt_todo)}")

    import concurrent.futures, threading
    lock = threading.Lock()
    stats = {"ok": 0, "fallback": 0, "error": 0}

    def _fix(i: int) -> None:
        tagged = para_tagged_text(paras[i])
        try:
            resp = client.chat.completions.create(
                model=MODEL,
                temperature=0.1,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": tagged},
                ],
            )
            fixed = resp.choices[0].message.content.strip()
            if ANY_BAD_RE.search(fixed) or not diff_is_safe(tagged, fixed):
                fixed = deterministic_fix(tagged)
                stats["fallback"] += 1
            else:
                stats["ok"] += 1
        except Exception as e:
            log.warning("абзац %d: %s — детерминированный фолбэк", i, e)
            fixed = deterministic_fix(tagged)
            stats["error"] += 1
        with lock:
            cache[str(i)] = fixed
            if len(cache) % 50 == 0:
                tmp = CACHE + ".tmp"
                json.dump(cache, open(tmp, "w", encoding="utf-8"), ensure_ascii=False)
                os.replace(tmp, CACHE)
                print(f"  прогресс: {len(cache)}/{len(affected)}")

    with concurrent.futures.ThreadPoolExecutor(max_workers=6) as ex:
        list(ex.map(_fix, gpt_todo))

    json.dump(cache, open(CACHE, "w", encoding="utf-8"), ensure_ascii=False)
    print(f"GPT: ok={stats['ok']}, fallback={stats['fallback']}, error={stats['error']}")

    # Применение к docx
    backup = f"{TARGET.rsplit('.', 1)[0]}_before_namefix_{datetime.now():%Y%m%d_%H%M%S}.docx"
    shutil.copy2(TARGET, backup)
    print(f"бэкап: {backup}")
    applied = 0
    for i_str, fixed in cache.items():
        i = int(i_str)
        if i < len(paras) and ANY_BAD_RE.search(paras[i].text):
            rewrite_para(paras[i], fixed)
            applied += 1
    doc.save(TARGET)
    print(f"применено: {applied} абзацев")

    # Контроль
    doc2 = Document(TARGET)
    left = sum(1 for p in doc2.paragraphs if ANY_BAD_RE.search(p.text))
    print(f"осталось абзацев с испорченными именами: {left}")
    if left == 0:
        os.remove(CACHE)
        print("Кэш удалён. Готово ✅")


if __name__ == "__main__":
    main()
