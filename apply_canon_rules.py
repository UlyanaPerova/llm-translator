#!/usr/bin/env python3
"""
Канон пользователя (2026-07-21), точечные правки без изменения прочего текста:
  1) Чёрный Плащ / Чёрная Накидка (все падежи) → Существо в чёрном плаще
     («Чёрный Плащ-ним» с хонорификом не трогаем)
  2) Куратор Ю (все падежи) → Ю Сонвун (падеж наследуется)
  3-4) «мистер + Имя» → только имя («мистер» без имени остаётся)
  5) Джованни → Джиованни
  6) одиночное «Ю» → «Ю Сонвун» в нужном падеже (GPT: падеж из контекста,
     пословный дифф-контроль, при сомнении не трогаем)
"""

import difflib
import json
import os
import re
import sys
import logging
import shutil
from datetime import datetime

from dotenv import load_dotenv
load_dotenv()

from logger import setup_logger
setup_logger(prefix="canon_rules")
log = logging.getLogger("canon_rules")

from docx import Document
from openai import OpenAI
from fix_name_poisoning import para_tagged_text, rewrite_para

TARGET = "The_Artist_Who_Paints_Dungeon_translated_pro.docx"
CACHE = ".canon_yu_cache.json"
MODEL = "gpt-5.1"

DETERMINISTIC = [
    # 2) Куратор Ю → Ю Сонвун с падежом (оба м.р., согласование не страдает)
    (r"[Кк]уратора\s+Ю\b(?![а-яё])", "Ю Сонвуна"),
    (r"[Кк]уратору\s+Ю\b(?![а-яё])", "Ю Сонвуну"),
    (r"[Кк]уратором\s+Ю\b(?![а-яё])", "Ю Сонвуном"),
    (r"[Кк]ураторе\s+Ю\b(?![а-яё])", "Ю Сонвуне"),
    (r"[Кк]уратор\s+Ю\b(?![а-яё])", "Ю Сонвун"),
    # 3-4) мистер + Имя → имя (мистер без имени остаётся)
    (r"[Мм]истер[а-яё]*\s+(?=[А-ЯЁ])", ""),
    # 5) Джованни → Джиованни
    (r"\bДжованни\b", "Джиованни"),
]

# 1) Плащ/Накидка → Существо в чёрном плаще: смена рода (м/ж → ср.) требует
# перенастройки согласования соседних слов — это работа GPT, не регулярки.
CLOAK_RE = re.compile(
    r"Ч[её]рн(?:ый|ого|ому|ым|ом)\s+Плащ[а-яё]*(?!-)|Ч[её]рн(?:ая|ой|ую)\s+Накидк[а-яё]*"
)
CLOAK_WORDS_RE = re.compile(r"^[«»\"'(\[]?(Ч[её]рн\w*|Плащ\w*|Накидк\w*)[.,!?;:…»\"')\]]*$")
CLOAK_INSERT_RE = re.compile(r"^[«»\"'(\[]?(Существ[оаеу]м?|в|ч[её]рном|плаще)[.,!?;:…»\"')\]]*$")

CLOAK_PROMPT = """You are a precise text-repair tool for a Russian novel translation.
The character name «Чёрный Плащ» / «Чёрная Накидка» (any grammatical form) must be
replaced everywhere with «Существо в чёрном плаще» in the correct grammatical case
(Существо/Существа/Существу/Существом/Существе в чёрном плаще).
The new head noun «Существо» is NEUTER — adjust the agreement of words that
grammatically depend on it (past-tense verbs, adjectives, participles):
«Чёрный Плащ кивнул» → «Существо в чёрном плаще кивнуло».
Do NOT touch «Чёрный Плащ-ним» (honorific form). Change nothing else.
Return ONLY the resulting paragraph text."""


def _strip_punct(w: str) -> str:
    return w.strip("«»\"'()[].,!?;:…")


def _is_agreement_pair(x: str, y: str) -> bool:
    """Перенастройка окончания того же слова: сказал→сказало, этот→это, был→было.
    Либо общий префикс ≥4, либо короткое слово — префикс длинного (разница ≤2)."""
    x, y = _strip_punct(x).lower(), _strip_punct(y).lower()
    if not x or not y:
        return False
    cp = os.path.commonprefix([x, y])
    if len(cp) >= 4:
        return True
    return len(cp) == min(len(x), len(y)) >= 2 and abs(len(x) - len(y)) <= 2


def cloak_diff_is_safe(original: str, fixed: str) -> bool:
    """Разрешены: удаление слов Плаща/Накидки, вставка слов новой фразы,
    и перенастройка окончаний у тех же основ (сказал→сказало, этот→это)."""
    a, b = original.split(), fixed.split()
    sm = difflib.SequenceMatcher(a=a, b=b, autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        removed = a[i1:i2]
        inserted = b[j1:j2]
        for w in inserted:
            if CLOAK_INSERT_RE.match(w):
                continue
            if any(_is_agreement_pair(w, r) for r in removed):
                continue
            return False
        for w in removed:
            if CLOAK_WORDS_RE.match(w):
                continue
            if any(_is_agreement_pair(w, ins) for ins in inserted):
                continue
            return False
    return True

BARE_YU_RE = re.compile(r"(?<![А-Яа-яЁё-])Ю(?![А-Яа-яё-])(?!\s+Сонвун)")
REMOVABLE_RE = re.compile(r"^[«»\"'(\[]?Ю[.,!?;:…»\"')\]]*$")
INSERTABLE_RE = re.compile(r"^[«»\"'(\[]?(Ю|Сонвун[а-яё]*)[.,!?;:…»\"')\]]*$")

YU_PROMPT = """You are a precise text-repair tool for a Russian novel translation.
In the paragraph, a standalone «Ю» is the bare surname of the character Ю Сонвун
(male). Korean etiquette forbids addressing by surname alone, so every standalone
«Ю» must become the full name «Ю Сонвун» in the grammatical case required by its
context: Ю Сонвун / Ю Сонвуна / Ю Сонвуну / Ю Сонвуном / Ю Сонвуне.
Do NOT touch occurrences already followed by «Сонвун», do NOT change any other
word or punctuation. Return ONLY the resulting paragraph text."""


def deterministic_pass(text: str) -> str:
    for pat, repl in DETERMINISTIC:
        text = re.sub(pat, repl, text)
    return text


def diff_is_safe(original: str, fixed: str) -> bool:
    a, b = original.split(), fixed.split()
    sm = difflib.SequenceMatcher(a=a, b=b, autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        for w in a[i1:i2]:
            if not REMOVABLE_RE.match(w):
                return False
        for w in b[j1:j2]:
            if not INSERTABLE_RE.match(w):
                return False
    return True


def main():
    api_key = os.getenv("OPENAI_API_KEY") or sys.exit("OPENAI_API_KEY не найден")
    client = OpenAI(api_key=api_key)

    doc = Document(TARGET)
    paras = doc.paragraphs

    backup = f"{TARGET.rsplit('.', 1)[0]}_before_canon_{datetime.now():%Y%m%d_%H%M%S}.docx"
    shutil.copy2(TARGET, backup)
    print(f"бэкап: {backup}")

    # Детерминированная фаза
    det_changed = 0
    for p in paras:
        tagged = para_tagged_text(p)
        fixed = deterministic_pass(tagged)
        if fixed != tagged:
            rewrite_para(p, fixed)
            det_changed += 1
    print(f"детерминированно изменено абзацев: {det_changed}")

    # GPT-фаза 1: Плащ/Накидка → Существо в чёрном плаще (со сменой согласования)
    cloak_affected = [i for i, p in enumerate(paras) if CLOAK_RE.search(p.text)]
    print(f"абзацев с Плащом/Накидкой: {len(cloak_affected)}")

    import concurrent.futures, threading
    cloak_lock = threading.Lock()
    cloak_stats = {"ok": 0, "guard": 0, "error": 0}
    cloak_results: dict[int, str] = {}

    def _fix_cloak(i: int) -> None:
        tagged = para_tagged_text(paras[i])
        try:
            resp = client.chat.completions.create(
                model=MODEL, temperature=0.1,
                messages=[{"role": "system", "content": CLOAK_PROMPT},
                          {"role": "user", "content": tagged}],
            )
            fixed = resp.choices[0].message.content.strip()
            if fixed != tagged and cloak_diff_is_safe(tagged, fixed) and not CLOAK_RE.search(fixed):
                cloak_stats["ok"] += 1
            else:
                fixed = tagged
                cloak_stats["guard"] += 1
        except Exception as e:
            log.warning("плащ, абзац %d: %s", i, e)
            fixed = tagged
            cloak_stats["error"] += 1
        with cloak_lock:
            cloak_results[i] = fixed
            if len(cloak_results) % 100 == 0:
                print(f"  Плащ-прогресс: {len(cloak_results)}/{len(cloak_affected)}")

    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as ex:
        list(ex.map(_fix_cloak, cloak_affected))
    for i, fixed in cloak_results.items():
        if para_tagged_text(paras[i]) != fixed:
            rewrite_para(paras[i], fixed)
    print(f"Плащ: ok={cloak_stats['ok']}, отклонено/оставлено={cloak_stats['guard']}, "
          f"ошибок={cloak_stats['error']}")

    # GPT-фаза 2: одиночное «Ю»
    affected = [i for i, p in enumerate(paras) if BARE_YU_RE.search(p.text)]
    print(f"абзацев с одиночным «Ю»: {len(affected)}")

    cache: dict[str, str] = {}
    if os.path.isfile(CACHE):
        cache = json.load(open(CACHE, encoding="utf-8"))

    import concurrent.futures, threading
    lock = threading.Lock()
    stats = {"ok": 0, "kept": 0, "guard": 0, "error": 0}

    def _fix(i: int) -> None:
        if str(i) in cache:
            return
        tagged = para_tagged_text(paras[i])
        try:
            resp = client.chat.completions.create(
                model=MODEL, temperature=0.1,
                messages=[{"role": "system", "content": YU_PROMPT},
                          {"role": "user", "content": tagged}],
            )
            fixed = resp.choices[0].message.content.strip()
            if fixed == tagged:
                stats["kept"] += 1
            elif diff_is_safe(tagged, fixed):
                stats["ok"] += 1
            else:
                fixed = tagged
                stats["guard"] += 1
        except Exception as e:
            log.warning("абзац %d: %s", i, e)
            fixed = tagged
            stats["error"] += 1
        with lock:
            cache[str(i)] = fixed
            if len(cache) % 100 == 0:
                tmp = CACHE + ".tmp"
                json.dump(cache, open(tmp, "w", encoding="utf-8"), ensure_ascii=False)
                os.replace(tmp, CACHE)
                print(f"  Ю-прогресс: {len(cache)}/{len(affected)}")

    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as ex:
        list(ex.map(_fix, affected))

    applied = 0
    for i_str, fixed in cache.items():
        i = int(i_str)
        if i < len(paras) and para_tagged_text(paras[i]) != fixed:
            rewrite_para(paras[i], fixed)
            applied += 1
    doc.save(TARGET)
    print(f"GPT: ok={stats['ok']}, без изменений={stats['kept']}, "
          f"отклонено диффом={stats['guard']}, ошибок={stats['error']}")
    print(f"применено Ю-правок: {applied}")

    # Контроль
    doc2 = Document(TARGET)
    text = "\n".join(p.text for p in doc2.paragraphs)
    checks = {
        "Чёрный Плащ (без -ним)": r"Ч[её]рн(?:ый|ого|ому|ым|ом)\s+Плащ[а-яё]*(?!-)",
        "Чёрная Накидка": r"Ч[её]рн(?:ая|ой|ую)\s+Накидк",
        "Куратор Ю": r"[Кк]уратор[а-яё]*\s+Ю\b",
        "мистер + Имя": r"[Мм]истер[а-яё]*\s+[А-ЯЁ][а-яё]+",
        "Джованни (старое)": r"\bДжованни\b",
        "одиночное Ю": BARE_YU_RE.pattern,
    }
    for name, pat in checks.items():
        print(f"осталось [{name}]: {len(re.findall(pat, text))}")
    if os.path.isfile(CACHE):
        os.remove(CACHE)


if __name__ == "__main__":
    main()
